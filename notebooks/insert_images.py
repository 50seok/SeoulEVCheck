"""
report_v2.docx 에 이미지 삽입
bottom-to-top 순서로 삽입 (인덱스 밀림 방지)
"""
import shutil
from pathlib import Path
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC  = Path(r'C:\teamwork\SeoulEVCheck\docs\report_v2.docx')
DEST = Path(r'C:\teamwork\SeoulEVCheck\docs\report_v3.docx')
NB   = Path(r'C:\teamwork\SeoulEVCheck\notebooks')

shutil.copy2(SRC, DEST)
doc = Document(DEST)

def insert_img_after(doc, body_idx, img_path, width=Inches(5.5)):
    """body[body_idx] 바로 뒤에 이미지 단락 삽입 (인라인 센터)"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(img_path), width=width)
    elem = p._element
    doc.element.body.remove(elem)          # 말미에서 꺼내서
    doc.element.body[body_idx].addnext(elem)  # 목표 위치 뒤에 삽입
    return elem

# ─── 아래→위 순서로 삽입 (인덱스 불변 유지) ──────────────────────────

# ⑥ [036] 활용 예시 뒤 → 예측 활용 코드
insert_img_after(doc, 36, NB / 'asset_code_predict.png', Inches(6.0))
print('[036] asset_code_predict.png 삽입')

# ⑤ [028] 베이스라인 초과 뒤 → 모델 비교표
insert_img_after(doc, 28, NB / 'model_compare_result.png', Inches(5.5))
print('[028] model_compare_result.png 삽입')

# ④ [026] 모델 정의 뒤 → 모델 학습 코드
insert_img_after(doc, 26, NB / 'asset_code_model.png', Inches(6.0))
print('[026] asset_code_model.png 삽입')

# ③ [017] 전처리 텍스트 뒤 → 전처리 코드
insert_img_after(doc, 17, NB / 'asset_code_preprocess.png', Inches(6.0))
print('[017] asset_code_preprocess.png 삽입')

# ② [015] 컬럼 설명 뒤 → 컬럼 표
insert_img_after(doc, 15, NB / 'asset_columns.png', Inches(6.2))
print('[015] asset_columns.png 삽입')

# ① [012] DATA IMPORTING 뒤 → 다이어그램
insert_img_after(doc, 12, NB / 'asset_diagram.png', Inches(6.2))
print('[012] asset_diagram.png 삽입')

doc.save(DEST)
print(f'\n저장 완료: {DEST}')
